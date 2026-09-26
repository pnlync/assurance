"""Build excel/single_policy_check.xlsx: an independent single-policy model in live Excel formulas. SPEC §13.4.

Every number on the model sheets is a formula over the Inputs sheet; only inputs (blue) and the Python values on
the Check sheets (the reconciliation targets) are typed in. Excel recalculates on open.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lifemodel.basis import load_basis
from lifemodel.curves import Curve
from lifemodel.parsers import PROCESSED
from lifemodel.portfolio import make_policy
from lifemodel.pricing import premium_from_rate, profit_test, solve_rate, solve_unisex_rate
from lifemodel.projection import project, value
from lifemodel.solvency2 import (CAT_SHOCK, EXPENSE_SHOCK, INFLATION_SHOCK, LAPSE_UP, MASS_LAPSE, MORTALITY_SHOCK,
                                 stress_losses)

OUT = Path(__file__).resolve().parents[2] / "excel" / "single_policy_check.xlsx"
VALUATION_RATE = 0.03  # flat curve used for V in both sheets (the fixture's valuation curve, SPEC §13.3)
FIRST, TERM = 8, 20     # first table row (t = 0) and policy term
LAST = FIRST + TERM     # row for t = 20

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF")
GREEN = Font(name="Arial", size=10, color="008000")
TITLE = Font(name="Arial", size=12, bold=True)
YELLOW = PatternFill("solid", fgColor="FFF2CC")

INPUT_ROWS = [  # key, label, note
    ("product", "Product", "LTA = level term; MP = mortgage protection"),
    ("age", "Issue age x", "Age last birthday at issue (SPEC §2.7)"),
    ("female", "Female (1/0)", ""),
    ("smoker", "Smoker (1/0)", ""),
    ("sa", "Sum assured / initial loan (€)", ""),
    ("loan_rate", "MP loan rate", "Used only for MP (SPEC §5)"),
    ("rate", "Premium rate per €1,000", "Solved in Python (brentq) for a 10% margin; Excel confirms the margin"),
    ("fee", "Policy fee (€ p.a.)", "SPEC §4.1"),
    ("mort_mult", "Mortality multiplier", "Fixture: 1. Reference: X = 0.985^22.5 (SPEC §4.1)"),
    ("base_q", "Fixture q at base age", "SPEC §13.3 (F1 only)"),
    ("base_age", "Fixture base age", "SPEC §13.3"),
    ("growth", "Fixture growth per year of age", "SPEC §13.3"),
    ("female_f", "Fixture female factor", "SPEC §13.3"),
    ("smoker_f", "Fixture smoker factor", "SPEC §13.3"),
    ("lapse1", "Lapse, policy year 1", "SPEC §4.1"),
    ("lapse2", "Lapse, policy year 2", ""),
    ("lapse3", "Lapse, policy year 3", ""),
    ("lapse4", "Lapse, policy year 4", ""),
    ("lapse5", "Lapse, policy year 5", ""),
    ("lapse_ult", "Lapse, year 6+ (0 in final year)", ""),
    ("lapse_mult", "Lapse multiplier", ""),
    ("acq", "Acquisition expense at t = 0 (€)", "SPEC §4.1"),
    ("maint", "Maintenance expense (€ p.a.)", ""),
    ("overhead", "Overhead (€ p.a.)", "Included here: Solvency II / pricing basis"),
    ("claim_exp", "Claim expense per death (€)", ""),
    ("infl", "Expense inflation", ""),
    ("comm", "Initial commission (share of premium)", ""),
    ("res_mort", "Reserving mortality multiplier", "SPEC §4.2: 110%, no lapses"),
    ("res_int", "Reserving interest", "SPEC §4.2"),
    ("earned", "Earned rate", "SPEC §4.3"),
    ("rdr", "Risk discount rate", "SPEC §4.3"),
    ("val_rate", "Valuation rate (flat) for V", "Flat 3%, as in the golden fixture"),
    ("term", "Term (years)", ""),
    ("sh_mort", "SII mortality shock", "Art. 137: +15%"),
    ("sh_up", "SII lapse-up factor", "Art. 142: × 1.5"),
    ("sh_mass", "SII mass lapse", "Art. 142: 40%"),
    ("sh_exp", "SII expense shock", "Art. 140: +10%"),
    ("sh_infl", "SII inflation shock", "Art. 140: +1pp"),
    ("sh_cat", "SII catastrophe shock", "Art. 143: +0.0015 for 12 months"),
]
ROW = {key: i + 3 for i, (key, _, _) in enumerate(INPUT_ROWS)}


def _inputs(basis, policy, fixture: bool) -> dict:
    m = basis.mortality_table
    lapses = list(basis.lapse_by_policy_year) + [basis.lapse_ultimate] * 5
    rate = (policy["premium"].iloc[0] - basis.policy_fee) * 1000 / policy["sa"].iloc[0]
    return {
        "product": policy["product"].iloc[0], "age": int(policy["issue_age"].iloc[0]),
        "female": int(policy["sex"].iloc[0] == "F"), "smoker": int(policy["smoker"].iloc[0]),
        "sa": float(policy["sa"].iloc[0]), "loan_rate": 0.04, "rate": float(rate), "fee": basis.policy_fee,
        "mort_mult": basis.mortality_multiplier,
        "base_q": m.base_q if fixture else None, "base_age": m.base_age if fixture else None,
        "growth": m.growth if fixture else None, "female_f": m.female_factor if fixture else None,
        "smoker_f": m.smoker_factor if fixture else None,
        **{f"lapse{i + 1}": lapses[i] for i in range(5)}, "lapse_ult": basis.lapse_ultimate,
        "lapse_mult": basis.lapse_multiplier, "acq": basis.acquisition, "maint": basis.maintenance,
        "overhead": basis.overhead, "claim_exp": basis.claim_expense, "infl": basis.inflation,
        "comm": basis.commission_first_year, "res_mort": basis.reserve_mortality_multiplier,
        "res_int": basis.reserve_interest, "earned": basis.earned_rate, "rdr": basis.risk_discount_rate,
        "val_rate": VALUATION_RATE, "term": TERM, "sh_mort": MORTALITY_SHOCK, "sh_up": LAPSE_UP,
        "sh_mass": MASS_LAPSE, "sh_exp": EXPENSE_SHOCK, "sh_infl": INFLATION_SHOCK, "sh_cat": CAT_SHOCK,
    }


def _model_sheet(wb, name: str, col: str, title: str, mortality: str):
    """One projection sheet. `col` is this policy's column on Inputs; `mortality` is the q formula template."""
    ws = wb.create_sheet(name)

    def i(key):
        return f"Inputs!${col}${ROW[key]}"

    ws["A1"], ws["A1"].font = title, TITLE
    summary = [
        ("B2", "Premium (€ p.a.)", "C2", f"={i('rate')}*{i('sa')}/1000+{i('fee')}"),
        ("B3", "NPV at RDR", "C3", f"=SUM(S{FIRST}:S{LAST - 1})"),
        ("B4", "EPV of premiums at RDR", "C4", f"=SUM(V{FIRST}:V{LAST - 1})"),
        ("B5", "Profit margin", "C5", "=C3/C4"),
        ("E2", "IRR", "F2", f"=IRR(Q{FIRST}:Q{LAST - 1})"),
        ("E3", "Discounted payback (years)", "F3", f"=MIN(U{FIRST}:U{LAST - 1})"),
        ("E4", "V at issue", "F4", f"=L{FIRST}"),
        ("H2", "SII loss: mortality", "I2", f"=MAX(0,W{FIRST}-L{FIRST})"),
        ("H3", "SII loss: lapse up", "I3", f"=IF(L{FIRST}<0,MAX(0,X{FIRST}-L{FIRST}),0)"),
        ("H4", "SII loss: mass lapse", "I4", f"={i('sh_mass')}*MAX(0,-L{FIRST})"),
        ("H5", "SII loss: expense", "I5", f"=MAX(0,Y{FIRST}-L{FIRST})"),
        ("H6", "SII loss: catastrophe", "I6",
         f"=MAX(0,K{FIRST}*{i('sh_cat')}*((F{FIRST}+J{FIRST})-(1-D{FIRST})*L{FIRST + 1}))"),
    ]
    for lab_cell, label, val_cell, formula in summary:
        ws[lab_cell], ws[lab_cell].font = label, FONT
        ws[val_cell], ws[val_cell].font = formula, BOLD
    ws["C5"].number_format = "0.0000%"
    ws["F2"].number_format = "0.0000%"

    headers = ["t", "Attained age", "q_t", "w_t", "ℓ_t (start of year)", "S_t", "Premium", "Commission",
               "Expenses (start)", "Claim expense", "v_t", "V_t best estimate", "q reserving", "V reserving",
               "Reserve _tV", "Profit vector Pr_t", "Profit signature Π_t", "RDR discount", "PV of Π_t",
               "Cumulative PV", "Payback helper", "PV premium term", "V mortality +15%", "V lapse up ×1.5",
               "V expense +10%, infl +1pp"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=7, column=j, value=h)
        c.font, c.alignment = BOLD, Alignment(wrap_text=True, vertical="bottom")
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.row_dimensions[7].height = 42

    for t in range(TERM + 1):
        r = FIRST + t
        nxt = r + 1
        cells = {"A": t, "B": f"={i('age')}+A{r}"}
        if t == TERM:  # maturity row: values and reserves are 0, ℓ continues
            cells.update({"E": f"=E{r - 1}*(1-C{r - 1})*(1-D{r - 1})", "L": 0, "N": 0, "O": 0, "W": 0, "X": 0, "Y": 0})
        else:
            cells.update({
                "C": mortality.format(r=r, i=i),
                "D": f"=IF(A{r}={i('term')}-1,0,MIN(1,{i('lapse_mult')}*IF(A{r}<5,INDEX(Inputs!${col}${ROW['lapse1']}"
                     f":${col}${ROW['lapse5']},A{r}+1),{i('lapse_ult')})))",
                "E": 1 if t == 0 else f"=E{r - 1}*(1-C{r - 1})*(1-D{r - 1})",
                "F": f"=IF({i('product')}=\"MP\",{i('sa')}*(1-(1/(1+{i('loan_rate')}))^({i('term')}-A{r}))"
                     f"/(1-(1/(1+{i('loan_rate')}))^{i('term')}),{i('sa')})",
                "G": "=$C$2",
                "H": f"=IF(A{r}=0,{i('comm')}*G{r},0)",
                "I": f"=IF(A{r}=0,{i('acq')},0)+({i('maint')}+{i('overhead')})*(1+{i('infl')})^A{r}",
                "J": f"={i('claim_exp')}*(1+{i('infl')})^A{r}",
                "K": f"=1/(1+{i('val_rate')})",
                "L": f"=-G{r}+H{r}+I{r}+K{r}*(C{r}*(F{r}+J{r})+(1-C{r})*(1-D{r})*L{nxt})",
                "M": f"=MIN(1,C{r}*{i('res_mort')})",
                "N": f"=-G{r}+H{r}+I{r}+(1/(1+{i('res_int')}))*(M{r}*(F{r}+J{r})+(1-M{r})*N{nxt})",
                "O": f"=IF(A{r}=0,0,MAX(0,N{r}))",
                "P": f"=(O{r}+G{r}-H{r}-I{r})*(1+{i('earned')})-C{r}*(F{r}+J{r})-(1-C{r})*(1-D{r})*O{nxt}",
                "Q": f"=E{r}*P{r}",
                "R": f"=(1+{i('rdr')})^-(A{r}+1)",
                "S": f"=Q{r}*R{r}",
                "T": f"=SUM(S${FIRST}:S{r})",
                "U": f"=IF(T{r}>=0,A{r}+1,999)",
                "V": f"=E{r}*G{r}*(1+{i('rdr')})^-A{r}",
                "W": f"=-G{r}+H{r}+I{r}+K{r}*(MIN(1,C{r}*(1+{i('sh_mort')}))*(F{r}+J{r})"
                     f"+(1-MIN(1,C{r}*(1+{i('sh_mort')})))*(1-D{r})*W{nxt})",
                "X": f"=-G{r}+H{r}+I{r}+K{r}*(C{r}*(F{r}+J{r})+(1-C{r})*(1-MIN(1,{i('sh_up')}*D{r}))*X{nxt})",
                "Y": f"=-G{r}+H{r}+(IF(A{r}=0,{i('acq')},0)+({i('maint')}+{i('overhead')})"
                     f"*(1+{i('infl')}+{i('sh_infl')})^A{r})*(1+{i('sh_exp')})+K{r}*(C{r}*(F{r}+{i('claim_exp')}"
                     f"*(1+{i('sh_exp')})*(1+{i('infl')}+{i('sh_infl')})^A{r})+(1-C{r})*(1-D{r})*Y{nxt})",
            })
        for letter, v in cells.items():
            c = ws[f"{letter}{r}"]
            c.value, c.font = v, FONT
            c.number_format = "0.00000000" if letter in "CDEKMR" else "#,##0.0000"
    ws.freeze_panes = "B8"
    ws["A30"] = ("Timing (SPEC §2.2): premium, commission and expenses at the start of each policy year; claims "
                 "and claim expenses at the end; lapses at the end, after deaths. V is computed backwards from V_20 = 0.")
    ws["A30"].font = Font(name="Arial", size=9, italic=True)
    return ws


def _check_sheet(wb, name: str, model: str, py: dict):
    """Python values (typed in) next to the Excel formulas, with differences."""
    ws = wb.create_sheet(name)
    ws["A1"], ws["A1"].font = f"Reconciliation: {model} vs the Python model (tolerance €0.01)", TITLE
    ws["A2"] = "Green = formula from the model sheet; black numbers = Python values; difference = Excel − Python."
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    rows = [("Premium", f"={model}!C2", py["premium"]), ("NPV", f"={model}!C3", py["npv"]),
            ("Profit margin", f"={model}!C5", py["margin"]), ("IRR", f"={model}!F2", py["irr"]),
            ("V at issue", f"={model}!F4", py["v0"])]
    rows += [(f"SII loss: {k}", f"={model}!I{n}", py["sii"][k])
             for n, k in zip(range(2, 7), ("mortality", "lapse_up", "mass", "expense", "cat"))]
    for j, h in enumerate(["Item", "Excel", "Python", "Difference"], start=1):
        ws.cell(row=4, column=j, value=h).font = BOLD
    for k, (label, formula, pyval) in enumerate(rows, start=5):
        ws.cell(row=k, column=1, value=label).font = FONT
        ws.cell(row=k, column=2, value=formula).font = GREEN
        ws.cell(row=k, column=3, value=float(pyval)).font = FONT
        ws.cell(row=k, column=4, value=f"=B{k}-C{k}").font = FONT
    top = 5 + len(rows) + 2
    cols = [("ℓ_t", "E", "l"), ("V_t", "L", "V"), ("Reserve", "O", "res"), ("Profit vector", "P", "pr")]
    ws.cell(row=top, column=1, value="t").font = BOLD
    for j, (label, _, _) in enumerate(cols):
        for k, suffix in enumerate(("Excel", "Python", "Diff")):
            ws.cell(row=top, column=2 + 3 * j + k, value=f"{label} {suffix}").font = BOLD
    for t in range(TERM):
        r = top + 1 + t
        ws.cell(row=r, column=1, value=t).font = FONT
        for j, (_, letter, key) in enumerate(cols):
            c0 = 2 + 3 * j
            ws.cell(row=r, column=c0, value=f"={model}!{letter}{FIRST + t}").font = GREEN
            ws.cell(row=r, column=c0 + 1, value=float(py[key][t])).font = FONT
            a, b = get_column_letter(c0), get_column_letter(c0 + 1)
            ws.cell(row=r, column=c0 + 2, value=f"={a}{r}-{b}{r}").font = FONT
    end = top + TERM
    ws.cell(row=end + 2, column=1, value="Largest absolute difference").font = BOLD
    diffs = [f"MAX(ABS(MIN(D5:D{4 + len(rows)})),ABS(MAX(D5:D{4 + len(rows)})))"]
    for j in range(len(cols)):
        d = get_column_letter(4 + 3 * j)
        diffs.append(f"MAX(ABS(MIN({d}{top + 1}:{d}{end})),ABS(MAX({d}{top + 1}:{d}{end})))")
    ws.cell(row=end + 2, column=2, value="=MAX(" + ",".join(diffs) + ")").font = BOLD
    ws.cell(row=end + 3, column=1, value="Reconciled within €0.01?").font = BOLD
    ws.cell(row=end + 3, column=2, value=f'=IF(B{end + 2}<0.01,"YES","NO")').font = BOLD
    for j in range(1, 14):
        ws.column_dimensions[get_column_letter(j)].width = 14
    return end + 2


def python_values(policy, basis) -> dict:
    """The Python model's numbers for one policy (flat 3% valuation curve)."""
    curve = Curve.flat(VALUATION_RATE)
    pt = profit_test(policy, basis)
    losses = stress_losses(policy, basis, curve)
    return {"premium": pt.premium, "npv": pt.npv, "margin": pt.margin, "irr": pt.irr,
            "v0": value(policy, basis, curve)[0, 0], "l": project(policy, basis)["in_force"][0],
            "V": value(policy, basis, curve)[0], "res": pt.reserves, "pr": pt.profit_vector,
            "sii": {k: losses[k][0, 0] for k in ("mortality", "lapse_up", "mass", "expense", "cat")}}


def build(path: Path = OUT) -> dict:
    """Write the workbook; returns the Python values and the cell holding each sheet's largest difference."""
    fixture, real = load_basis("fixture"), load_basis("basis_2022")
    f1 = make_policy("LTA", 35, "M", False, 250_000,
                     premium_from_rate(solve_rate("LTA", 35, "M", False, fixture), 250_000, fixture))
    ref = make_policy("LTA", 35, "M", False, 250_000,
                      premium_from_rate(solve_unisex_rate("LTA", 35, False, real), 250_000, real))

    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"], ws["A1"].font = "Inputs (blue = typed-in assumption; everything else in the workbook is a formula)", TITLE
    for j, h in enumerate(["Assumption", "F1 (golden fixture)", "Reference (real basis)", "Source / note"], start=1):
        ws.cell(row=2, column=j, value=h).font = BOLD
    for col, values in (("B", _inputs(fixture, f1, True)), ("C", _inputs(real, ref, False))):
        for key, label, note in INPUT_ROWS:
            ws[f"A{ROW[key]}"], ws[f"A{ROW[key]}"].font = label, FONT
            ws[f"D{ROW[key]}"], ws[f"D{ROW[key]}"].font = note, Font(name="Arial", size=9, italic=True)
            c = ws[f"{col}{ROW[key]}"]
            c.value, c.font = values[key], BLUE
            if key == "rate":
                c.fill = YELLOW
                c.comment = Comment("Solved in Python for a 10% margin. In Excel: Data > What-If > Goal Seek, "
                                    "set the model sheet's C5 to 10% by changing this cell.", "lifemodel")
    ws["C" + str(ROW["base_q"])] = "n/a (CMI 00 table)"
    for letter, width in (("A", 36), ("B", 20), ("C", 22), ("D", 60)):
        ws.column_dimensions[letter].width = width

    fixture_q = ("=MIN(1,{i}*{bq}*{g}^(B{{r}}-{ba})*IF({f}=1,{ff},1)*IF({s}=1,{sf},1))")
    q_f1 = fixture_q.format(i="Inputs!$B$" + str(ROW["mort_mult"]), bq="Inputs!$B$" + str(ROW["base_q"]),
                            g="Inputs!$B$" + str(ROW["growth"]), ba="Inputs!$B$" + str(ROW["base_age"]),
                            f="Inputs!$B$" + str(ROW["female"]), ff="Inputs!$B$" + str(ROW["female_f"]),
                            s="Inputs!$B$" + str(ROW["smoker"]), sf="Inputs!$B$" + str(ROW["smoker_f"]))
    # Reference life: male non-smoker → TMN00, select by duration (0–4, then 5+)
    q_ref = ("=MIN(1,Inputs!$C$" + str(ROW["mort_mult"]) + "*INDEX(CMI_TMN00!$B$3:$G$106,"
             "MATCH(B{r},CMI_TMN00!$A$3:$A$106,0),MIN(A{r},5)+1))")
    _model_sheet(wb, "F1", "B", "F1 — level term, male 35 non-smoker, €250k, 20 years (golden fixture basis)",
                 q_f1.replace("{{r}}", "{r}"))
    _model_sheet(wb, "Reference", "C",
                 "Reference — level term, male 35 non-smoker, €250k, 20 years (real basis: CMI 00 × X, unisex rate)",
                 q_ref)

    cmi = pd.read_csv(PROCESSED / "mortality_cmi00.csv")
    tmn = cmi[cmi["table"] == "TMN00"].pivot(index="age", columns="duration", values="qx")
    ws = wb.create_sheet("CMI_TMN00")
    ws["A1"] = "CMI 00 series, TMN00 (temporary assurances, male non-smokers), q[x−t]+t by attained age. Source: IFoA."
    ws["A1"].font = BOLD
    for j, h in enumerate(["Attained age", "Dur 0", "Dur 1", "Dur 2", "Dur 3", "Dur 4", "Dur 5+"], start=1):
        ws.cell(row=2, column=j, value=h).font = BOLD
    for k, (age, row) in enumerate(tmn.iterrows(), start=3):
        ws.cell(row=k, column=1, value=int(age)).font = FONT
        for d in range(6):
            v = row.get(d)
            ws.cell(row=k, column=2 + d, value=None if pd.isna(v) else float(v)).font = BLUE

    py_f1, py_ref = python_values(f1, fixture), python_values(ref, real)
    cell_f1 = _check_sheet(wb, "Check_F1", "F1", py_f1)
    cell_ref = _check_sheet(wb, "Check_Reference", "Reference", py_ref)
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"F1": py_f1, "Reference": py_ref, "max_diff_cells": {"Check_F1": f"B{cell_f1}", "Check_Reference": f"B{cell_ref}"}}


if __name__ == "__main__":
    info = build()
    print(OUT, info["max_diff_cells"])
